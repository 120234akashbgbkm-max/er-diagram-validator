"""
file_ingest.py
Turns ANY uploaded file into the model that er_validator.validate() expects.

  - draw.io files with table shapes -> parsed exactly (er_parser), no AI
  - .sql / text containing CREATE TABLE -> parsed exactly (sql_parser), no AI
  - everything else -> read (PDF / images sent to Gemini directly, Word / Office / text
    converted to text) and Gemini extracts tables, columns and relationships.

The AI only READS the file; the rule-based validator judges the result.
"""

import io
import os
import re
import zipfile

from google.genai import types

from er_model import build_model
from er_parser import drawio_xml_text, parse_drawio
from llm import generate_json
from sql_parser import parse_sql

MAX_BYTES = 15 * 1024 * 1024      # 15 MB upload limit
MAX_TEXT_CHARS = 150_000          # text sent to the AI
MAX_IMAGES = 8                    # images pulled out of Word/Office files

MAX_AI_PAYLOAD = 14 * 1024 * 1024  # Gemini inline requests cap near 20 MB once base64-encoded

IMAGE_MIMES = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".webp": "image/webp", ".heic": "image/heic", ".heif": "image/heif",
}
CONVERT_TO_PNG = {".gif", ".bmp", ".tif", ".tiff", ".ico", ".jfif"}
VECTOR_PICS = {".emf", ".wmf"}      # Word pastes drawings like this; readable only where Pillow supports them
_SKIPPED = {"pictures": 0}

EXTRACT_PROMPT = """You are reading an ER (entity-relationship) diagram or database
design from an uploaded file. The file may be a picture, a PDF, a Word document,
SQL DDL, JSON, XML, PlantUML/Mermaid/DBML, plain text, or a written description.

Extract EVERY table/entity, its columns and its relationships, exactly as drawn or
written. Do not invent tables or columns that are not in the file, and do not fix
mistakes in the design - the diagram will be validated afterwards, so copy its
problems faithfully (missing keys, missing types, duplicate names, etc.).

Return ONLY a JSON object:
{
  "diagram_found": true or false,
  "notation": "chen" or "crowsfoot" or "other",
  "entities": [
    {"name": "Orders",
     "attributes": [
       {"name": "order_id", "type": "INT", "is_pk": true, "is_fk": false, "is_unique": false}
     ]}
  ],
  "relationships": [
    {"name": "places", "from": "Customers", "to": "Orders",
     "from_cardinality": "1", "to_cardinality": "0..N"}
  ]
}

Rules:
- "notation": "chen" for rectangles + ovals + diamonds, "crowsfoot" for tables with a row per
  column, "other" for anything else (SQL, text, UML, Mermaid...). Never invent a notation.
- "from" / "to" must be entity names exactly as you wrote them in "entities".
- Cardinality values must be one of "1", "0..1", "N", "1..N", "0..N", or null when the
  file does not show it. Cardinality is the count on THAT end of the line
  (from_cardinality = how many "from" rows per related row).
- "type" is "" when the file shows no data type.
- Mark is_pk / is_fk / is_unique from PK/FK/UK labels, underlining, REFERENCES,
  PRIMARY KEY, FOREIGN KEY or UNIQUE clauses.
- KEYS IN PICTURES: look closely at every attribute label. In Chen / oval-style
  diagrams an UNDERLINED attribute name is a primary key (is_pk = true). A dashed or
  dotted underline marks the partial key of a weak entity - also set is_pk = true
  for it. Check each entity one by one before answering; do not skip underlines.
- For Chen-style diagrams, treat each entity as a table and its attributes as columns.
  A diamond is a relationship; double-line boxes/diamonds are weak entities and
  identifying relationships. Multi-valued (double oval) and derived (dashed oval)
  attributes are still listed as normal columns.
- If the file contains no ER diagram or database design at all, return
  {"diagram_found": false, "entities": [], "relationships": []}."""


# ---------------------------------------------------------------- detection

def _looks_like_text(raw):
    """Return decoded text if the bytes look like human-readable text."""
    sample = raw[:20000]
    if b"\x00" in sample and not (sample[:2] in (b"\xff\xfe", b"\xfe\xff")):
        return None
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
        if not text.strip():
            return None
        printable = sum(ch.isprintable() or ch in "\r\n\t" for ch in text[:20000])
        if printable / max(1, len(text[:20000])) > 0.92:
            return text
    return None


def _detect(filename, raw):
    ext = os.path.splitext(filename or "")[1].lower()
    head = raw[:16]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\x89PNG"):
        return "image"
    if head.startswith(b"\xff\xd8\xff"):
        return "image"
    if head.startswith((b"GIF87a", b"GIF89a")) or (head[:4] == b"RIFF" and raw[8:12] == b"WEBP"):
        return "image"
    if head.startswith(b"PK\x03\x04"):
        return "office_zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "legacy_office"
    if ext in IMAGE_MIMES or ext in CONVERT_TO_PNG:
        return "image"
    if ext == ".pdf":
        return "pdf"
    return "text_or_binary"


MAX_SIDE = 3000          # longest side sent to the AI (plenty to read small ER text)
MAX_IMAGE_BYTES = 8 * 1024 * 1024


def _is_heic(raw, ext):
    return ext in (".heic", ".heif") or raw[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1", b"ftypheif", b"ftyphevc")


def _image_part(raw, ext_hint=""):
    """Clean ANY image so the AI can read it, then wrap it as a Gemini part.

    Fixes the things that make pictures unreadable: phone photos stored sideways (EXIF
    rotation), transparent backgrounds (black lines vanish on black), CMYK / 16-bit / palette
    modes, animated GIFs (first frame), and huge dimensions.
    """
    ext = ext_hint.lower()
    try:
        from PIL import Image, ImageOps
    except ImportError:
        raise ValueError("Reading images needs the Pillow package: pip install Pillow")

    try:
        img = Image.open(io.BytesIO(raw))
        was_jpeg = img.format == "JPEG"
        img.seek(0)                                  # animated GIF / multi-page TIFF: first frame
        img = ImageOps.exif_transpose(img)           # upright phone photos

        if img.mode in ("I;16", "I;16L", "I;16B", "I"):  # 16 / 32-bit grayscale -> 8-bit
            img = img.point(lambda v: v / 256).convert("L")
        has_alpha = img.mode in ("RGBA", "LA", "PA") or (img.mode == "P" and "transparency" in img.info)
        if has_alpha:                                # flatten transparency onto white
            rgba = img.convert("RGBA")
            flat = Image.new("RGB", rgba.size, "white")
            flat.paste(rgba, mask=rgba.split()[3])
            img = flat
        elif img.mode != "RGB":                      # CMYK, grayscale, palette, 1-bit
            img = img.convert("RGB")

        if max(img.size) > MAX_SIDE:                 # shrink huge images
            img.thumbnail((MAX_SIDE, MAX_SIDE), Image.LANCZOS)

        buf = io.BytesIO()
        if was_jpeg:
            img.save(buf, "JPEG", quality=92)
            mime = "image/jpeg"
        else:
            img.save(buf, "PNG")
            mime = "image/png"
            if buf.tell() > MAX_IMAGE_BYTES:         # photo-like PNG: JPEG is far smaller
                buf = io.BytesIO()
                img.save(buf, "JPEG", quality=88)
                mime = "image/jpeg"
        return types.Part.from_bytes(data=buf.getvalue(), mime_type=mime)
    except Exception:
        # Pillow cannot decode HEIC without a plugin, but Gemini reads it natively
        if _is_heic(raw, ext):
            return types.Part.from_bytes(data=raw, mime_type=IMAGE_MIMES.get(ext, "image/heic"))
        raise ValueError(
            "Could not read this image file. It may be damaged or in an unsupported format. "
            "Try saving it as PNG or JPG."
        )


# ------------------------------------------------------- Office-style files

def _strip_xml(xml_bytes):
    text = xml_bytes.decode("utf-8", errors="replace")
    text = re.sub(r"</(w:p|a:p|text:p|row|si)>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"[ \t]+", " ", text).strip()


def _read_docx_text(raw):
    """Paragraphs and tables of a .docx, in a simple readable form."""
    try:
        from docx import Document
    except ImportError:
        return None
    doc = Document(io.BytesIO(raw))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for ti, table in enumerate(doc.tables, 1):
        lines.append(f"\n[Table {ti}]")
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))

    # Diagrams drawn with Word shapes keep their text in text boxes, which python-docx skips
    try:
        xml = zipfile.ZipFile(io.BytesIO(raw)).read("word/document.xml").decode("utf-8", "replace")
        boxes = [_strip_xml(m.encode()) for m in re.findall(r"<w:txbxContent>(.*?)</w:txbxContent>", xml, re.S)]
        boxes = list(dict.fromkeys(b for b in boxes if b))   # Word stores each box twice; keep one
        if boxes:
            lines.append("\n[Text inside shapes and text boxes]")
            lines.extend(boxes)
    except Exception:
        pass
    return "\n".join(lines)


def _read_office_zip(raw):
    """Returns (text, [image parts]) for docx / xlsx / pptx / odt / vsdx etc."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        raise ValueError("This archive is damaged or is not a supported document.")

    names = zf.namelist()
    text_parts = []

    if "word/document.xml" in names:
        text = None
        try:
            text = _read_docx_text(raw)
        except Exception:
            text = None
        if not text:
            text = _strip_xml(zf.read("word/document.xml"))
        text_parts.append(text)
    no_shared_strings = "xl/workbook.xml" in names and "xl/sharedStrings.xml" not in names
    for n in names:
        if no_shared_strings and n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
            text_parts.append(_strip_xml(zf.read(n)))
            continue
        if (n == "xl/sharedStrings.xml" or n.startswith("ppt/slides/slide")
                or n == "content.xml" or n.startswith("visio/pages/page")) and n.endswith(".xml"):
            text_parts.append(_strip_xml(zf.read(n)))

    images = []
    for n in names:
        ext = os.path.splitext(n)[1].lower()
        if "/media/" in n and (ext in IMAGE_MIMES or ext in CONVERT_TO_PNG or ext in VECTOR_PICS):
            try:
                images.append(_image_part(zf.read(n), ext))
            except ValueError:
                _SKIPPED["pictures"] += 1
                continue
            if len(images) >= MAX_IMAGES:
                break

    # a zip of plain text files (e.g. a project folder): include readable ones
    if not text_parts and not images:
        for n in names:
            if n.endswith("/") or "__MACOSX" in n:
                continue
            data = zf.read(n)
            if len(data) > 2_000_000:
                continue
            t = _looks_like_text(data)
            if t:
                text_parts.append(f"--- {n} ---\n{t}")

    return "\n\n".join(p for p in text_parts if p), images


def _printable_strings(raw):
    """Last-resort text recovery from legacy binaries (.doc etc.)."""
    runs = re.findall(rb"[\x20-\x7e\r\n\t]{5,}", raw)
    runs += [r.decode("utf-16le", "ignore").encode() for r in
             re.findall(rb"(?:[\x20-\x7e]\x00){5,}", raw)]
    return "\n".join(r.decode("utf-8", "ignore") for r in runs)


# ------------------------------------------------------------ AI extraction

def _normalize(data):
    """Gemini output -> the model shape er_parser produces (ids, layout, cardinalities)."""
    tables = []
    for e in data.get("entities") or []:
        if not isinstance(e, dict):
            continue
        cols = []
        for a in e.get("attributes") or []:
            if isinstance(a, str):
                a = {"name": a}
            if not isinstance(a, dict) or not str(a.get("name") or "").strip():
                continue
            keys = [k for k, flag in (("PK", "is_pk"), ("FK", "is_fk"), ("UK", "is_unique")) if a.get(flag)]
            cols.append({"name": a["name"], "type": a.get("type") or "", "keys": keys})
        tables.append({"name": e.get("name") or "", "columns": cols})

    rels = [r for r in (data.get("relationships") or []) if isinstance(r, dict)]
    notation = data.get("notation") if data.get("notation") in ("chen", "crowsfoot") else "other"
    return build_model({"notation": notation, "tables": tables, "relationships": rels})


def _payload_size(parts):
    total = 0
    for p in parts:
        if isinstance(p, str):
            total += len(p.encode("utf-8"))
        else:
            inline = getattr(p, "inline_data", None)
            total += len(getattr(inline, "data", b"") or b"")
    return total


def _extract_with_ai(parts):
    if not os.getenv("GEMINI_API_KEY"):
        raise RuntimeError(
            "GEMINI_API_KEY is not set. It is needed to read files that are not draw.io or SQL. "
            "Check your .env file."
        )
    if _payload_size(parts) > MAX_AI_PAYLOAD:
        raise ValueError(
            "This file is too large for the AI to read (limit about 14 MB). "
            "Compress it, or export just the ER diagram as a smaller PNG or PDF."
        )

    # retries on 503 and falls back to other models (see llm.py)
    data = generate_json(EXTRACT_PROMPT, contents=parts, temperature=0.0)

    if isinstance(data, list):  # model returned a bare list of entities
        data = {"entities": data, "relationships": []}
    if not isinstance(data, dict) or not data.get("diagram_found", True) or not data.get("entities"):
        raise ValueError(
            "No ER diagram or database tables could be found in this file. "
            "Upload a file that contains an ER diagram, a schema or a table design."
        )
    model = _normalize(data)
    if not model["entities"]:
        raise ValueError(
            "No ER diagram or database tables could be found in this file. "
            "Upload a file that contains an ER diagram, a schema or a table design."
        )
    return model


# -------------------------------------------------------------- entry point

def ingest(filename, raw):
    """Returns (model, info) where info = {"source", "file_type"}.

    source is one of: drawio, sql, ai-drawio, ai-image, ai-pdf, ai-office, ai-text.
    """
    if not raw:
        raise ValueError("The uploaded file is empty.")
    if len(raw) > MAX_BYTES:
        raise ValueError("File is too large. The limit is 15 MB.")

    ext = os.path.splitext(filename or "")[1].lower()
    kind = _detect(filename, raw)
    label = ext.lstrip(".").upper() or "file"
    ai_source = None

    parts = [
        "Extract the ER diagram / database design from the following "
        f"uploaded file (name: {filename or 'unknown'})."
    ]

    # 1) draw.io / mxGraph files: exact parser when the diagram uses table shapes
    if b"<mxfile" in raw[:4000] or b"<mxGraphModel" in raw[:4000]:
        try:
            return parse_drawio(raw), {"source": "drawio", "file_type": label}
        except Exception:
            pass  # not table shapes (e.g. Chen notation): let the AI read the decompressed XML
        try:
            xml = drawio_xml_text(raw)
        except Exception:
            raise ValueError("This draw.io file could not be read. Try exporting it again as uncompressed XML.")
        parts.append(
            "This is the XML of a draw.io ER diagram. Each <mxCell vertex=\"1\"> is a shape: its text is "
            "'value' and 'style' gives the type (ellipse = attribute, rhombus = relationship, plain "
            "rectangle = entity, swimlane = table). Cells with edge=\"1\" join the shapes named in "
            "'source' and 'target'; an edge's value is usually a cardinality or relationship name. "
            "Ignore coordinates.\n\nXML:\n" + xml[:MAX_TEXT_CHARS]
        )
        ai_source = "ai-drawio"

    elif kind == "pdf":
        parts.append(types.Part.from_bytes(data=raw, mime_type="application/pdf"))
        ai_source = "ai-pdf"
    elif kind == "image":
        parts.append(_image_part(raw, ext))
        ai_source = "ai-image"
    elif kind == "office_zip":
        _SKIPPED["pictures"] = 0
        text, images = _read_office_zip(raw)
        if not text and not images:
            hint = ""
            if _SKIPPED["pictures"]:
                hint = (f" {_SKIPPED['pictures']} picture(s) inside it could not be decoded (Word often stores "
                        "pasted drawings in a format that cannot be read). Save the document as PDF and upload that instead.")
            raise ValueError("Could not find any readable content in this document." + hint)
        if text:
            parts.append("FILE TEXT:\n" + text[:MAX_TEXT_CHARS])
        parts.extend(images)
        ai_source = "ai-office"
    elif kind == "legacy_office":
        text = _printable_strings(raw)
        if len(text) < 30:
            raise ValueError(
                "This old binary Office format could not be read. "
                "Please save it as .docx or PDF and upload again."
            )
        parts.append("FILE TEXT (recovered from a binary file):\n" + text[:MAX_TEXT_CHARS])
        ai_source = "ai-office"
    else:
        text = _looks_like_text(raw)
        if text is None:
            raise ValueError(
                "This file type could not be read. Try PDF, an image (PNG/JPG), "
                "Word (.docx), draw.io, or a text/SQL file."
            )
        # 2) SQL DDL: parse exactly, no AI. Falls through to the AI if it is not really SQL.
        if re.search(r"CREATE\s+(?:\w+\s+)*TABLE", text, re.I):
            try:
                return build_model(parse_sql(text)), {"source": "sql", "file_type": label}
            except Exception:
                pass
        parts.append("FILE TEXT:\n" + text[:MAX_TEXT_CHARS])
        ai_source = "ai-text"

    try:
        model = _extract_with_ai(parts)
    except ValueError as e:
        if ai_source == "ai-office" and _SKIPPED["pictures"]:
            raise ValueError(
                f"{e} Note: {_SKIPPED['pictures']} picture(s) inside the document could not be read "
                "(Word often stores pasted drawings in a format that cannot be decoded). "
                "Save the document as PDF and upload that instead."
            )
        raise
    return model, {"source": ai_source, "file_type": label}
